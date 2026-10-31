"""v1.7 工科正式版产物校验（测试用途 · 虚构数据）：pdftotext 可检索与阅读顺序、pdffonts 字体嵌入、
python-docx 解析、LibreOffice headless DOCX→PDF→PNG。用法：python3 tests/e2e/presets_verify.py <artifacts_dir>"""
import json, subprocess, sys, os, glob, shutil
from docx import Document

D = sys.argv[1] if len(sys.argv) > 1 else '/workspace/resume-tool/artifacts/p12'
EXPECT = {
    'CN': ['张测试', '教育背景', '项目经历', '实习 / 工作经历', '专业技能', '校园经历', '证书及奖项'],
    'EN': ['Ce Zhang', 'EDUCATION', 'PROJECT EXPERIENCE', 'PROFESSIONAL EXPERIENCE', 'TECHNICAL SKILLS', 'LEADERSHIP & ACTIVITIES', 'CERTIFICATIONS AND AWARDS'],
}
MUST = {'CN': ['2025.09 – 2028.03（预计）', 'Φ90H7', 'Ra 1.6 μm', '≤ 0.02 mm', 'GD&T', 'MATLAB/Simulink', '课程设计'],
        'EN': ['Sep 2025 – Mar 2028 (expected)', 'Φ90H7', 'Ra 1.6 μm', '≤ 0.02 mm', 'Dassault Systèmes', 'Course Project', 'GD&T', 'English CET-6 (score 532)']}
FORBID = ['出生年月', '性别', '籍贯', '政治面貌', 'Gender', 'Hometown', 'Born', 'English: CET-6', '\ufffd']
HEADS = {'教育背景', '项目经历', '实习 / 工作经历', '专业技能', '校园经历', '证书及奖项', 'EDUCATION', 'PROJECT EXPERIENCE', 'PROFESSIONAL EXPERIENCE', 'TECHNICAL SKILLS', 'LEADERSHIP & ACTIVITIES', 'CERTIFICATIONS AND AWARDS'}

def page_checks(pdf, label):
    """逐页：页尾不能是孤立的模块标题；不能有替换字符（乱码）"""
    info = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
    pages = int([l for l in info.splitlines() if l.startswith('Pages')][0].split()[-1])
    orphan = []
    for i in range(1, pages + 1):
        t = subprocess.run(['pdftotext', '-f', str(i), '-l', str(i), '-layout', pdf, '-'], capture_output=True, text=True).stdout
        lines = [l.strip() for l in t.splitlines() if l.strip()]
        if lines and lines[-1] in HEADS: orphan.append(i)
        check('\ufffd' not in t, f'{label}: 第 {i} 页有乱码')
    check(not orphan, f'{label}: 第 {orphan} 页页尾孤立标题')
    return pages

def campus_ok(n, lang, label):
    """社团经历只能出现在校园 / LEADERSHIP 模块"""
    key = '机械创新协会' if lang == 'CN' else 'Mechanical Innovation Society'
    work_h, next_h, camp_h = ('实习 / 工作经历', '专业技能', '校园经历') if lang == 'CN' else ('PROFESSIONAL EXPERIENCE', 'TECHNICAL SKILLS', 'LEADERSHIP & ACTIVITIES')
    w0, w1, c0, k = n.find(work_h), n.find(next_h), n.find(camp_h), n.find(key)
    return check(k > c0 >= 0 and not (w0 <= k < w1), f'{label}: 社团经历位置错误')
res, fails = {}, []

def check(cond, msg):
    if not cond: fails.append(msg)
    return cond

def norm(t): return ' '.join(t.split())

for pdf in sorted(glob.glob(os.path.join(D, '*_CN_*.pdf')) + glob.glob(os.path.join(D, '*_EN_*.pdf'))):
    name = os.path.basename(pdf); lang = 'CN' if '_CN_' in name else 'EN'
    txt = subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True).stdout
    raw = subprocess.run(['pdftotext', pdf, '-'], capture_output=True, text=True).stdout
    n = norm(raw)
    fonts = subprocess.run(['pdffonts', pdf], capture_output=True, text=True).stdout.splitlines()[2:]
    pos = [n.find(h) for h in EXPECT[lang]]
    r = {'chars': len(n), 'headings_found': all(p >= 0 for p in pos), 'reading_order_ok': pos == sorted(pos) and all(p >= 0 for p in pos),
         'must_present': {m: (m in n) for m in MUST[lang]}, 'forbidden_absent': all(f not in n for f in FORBID),
         'fonts_all_embedded': all(' yes ' in f for f in fonts), 'font_names': sorted({f.split()[0].split('+')[-1] for f in fonts})}
    check(r['reading_order_ok'], f'{name}: 标题顺序 {pos}')
    check(all(r['must_present'].values()), f'{name}: 缺字/缺文本 {[k for k, v in r["must_present"].items() if not v]}')
    check(r['forbidden_absent'], f'{name}: 出现个人敏感字段')
    check(r['fonts_all_embedded'], f'{name}: 有未嵌入字体')
    r['pages'] = page_checks(pdf, name); r['campus_ok'] = campus_ok(n, lang, name)
    res[name] = r

# DOCX：python-docx 解析 + LibreOffice 渲染
lo_out = os.path.join(D, 'docx_render'); os.makedirs(lo_out, exist_ok=True)
for dx in sorted(glob.glob(os.path.join(D, '*_CN_*.docx')) + glob.glob(os.path.join(D, '*_EN_*.docx'))):
    name = os.path.basename(dx); lang = 'CN' if '_CN_' in name else 'EN'
    doc = Document(dx)
    paras = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells: paras += [p.text for p in c.paragraphs]
    full = norm('\n'.join(paras))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    fonts = set()
    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.name: fonts.add(run.font.name)
    st = doc.styles['Normal'].font
    rFonts = doc.styles['Normal'].element.rPr.rFonts if doc.styles['Normal'].element.rPr is not None else None
    ea = rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia') if rFonts is not None else None
    images = len(doc.inline_shapes)
    import zipfile
    xml = zipfile.ZipFile(dx).read('word/document.xml').decode()
    check('txbxContent' not in xml and '<w:pict' not in xml, f'{name}: 使用了文本框')
    if images:
        t0 = doc.tables[0]; last = t0.rows[0].cells[-1]
        check('graphic' in last._tc.xml and 'wp:inline' in last._tc.xml, f'{name}: 照片不在页眉右侧单元格')
    r = {'paragraphs': len(doc.paragraphs), 'headings': heads, 'campus_ok': campus_ok(full, lang, name + ' docx'), 'normal_font': st.name, 'eastAsia': ea, 'run_font_overrides': sorted(fonts), 'images': images,
         'must_present': {m: (m in full) for m in MUST[lang]}}
    check(heads == [h for h in EXPECT[lang][1:] if h in heads] and len(heads) >= 5, f'{name}: 标题 {heads}')
    check(all(r['must_present'].values()), f'{name}: DOCX 缺文本 {[k for k, v in r["must_present"].items() if not v]}')
    check(images == (1 if 'Photo' in name or '有照片' in name else 0) or 'NoPhoto' in name, f'{name}: 图片数 {images}')
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', lo_out, dx], capture_output=True, timeout=180)
    lpdf = os.path.join(lo_out, name[:-5] + '.pdf')
    if check(os.path.exists(lpdf), f'{name}: LibreOffice 转换失败'):
        pages = int([l for l in subprocess.run(['pdfinfo', lpdf], capture_output=True, text=True).stdout.splitlines() if l.startswith('Pages')][0].split()[-1])
        r['lo_pages'] = pages
        page_checks(lpdf, name + ' (LO)')
        check(pages == res.get(name[:-5] + '.pdf', {}).get('pages', pages), f'{name}: LO 页数 {pages} 与 PDF 不同')
        subprocess.run(['pdftoppm', '-r', '90', '-png', lpdf, os.path.join(D, 'render_docx_' + name[:-5])])
        lt = norm(subprocess.run(['pdftotext', lpdf, '-'], capture_output=True, text=True).stdout)
        r['lo_text_has_headings'] = all(h in lt for h in EXPECT[lang][1:])
        check(r['lo_text_has_headings'], f'{name}: LO 渲染后标题缺失')
    res[name] = r

# 边界样例：逐页孤立标题 / 乱码；DOCX 同样经 LibreOffice 渲染
for pdf in sorted(glob.glob(os.path.join(D, 'edge_*.pdf'))):
    res[os.path.basename(pdf)] = {'pages': page_checks(pdf, os.path.basename(pdf))}
for dx in sorted(glob.glob(os.path.join(D, 'edge_*.docx'))):
    name = os.path.basename(dx); Document(dx)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', lo_out, dx], capture_output=True, timeout=180)
    lpdf = os.path.join(lo_out, name[:-5] + '.pdf')
    if check(os.path.exists(lpdf), f'{name}: LibreOffice 转换失败'):
        res[name] = {'lo_pages': page_checks(lpdf, name + ' (LO)')}
        subprocess.run(['pdftoppm', '-r', '70', '-png', lpdf, os.path.join(D, 'render_docx_' + name[:-5])])

# 学位状态：四套预设 PDF 与 DOCX 的学历日期文字逐条一致
import re
DEG = {'zh': ['2024.09 – 2026.12（预计 2027.01 授予学位）', '2020.09 – 2024.06（预计）', '2016.09 – 2020.06'],
       'en': ['Sep 2024 – Dec 2026 (degree expected Jan 2027)', 'Sep 2020 – Jun 2024 (expected)', 'Sep 2016 – Jun 2020']}
for pdf in sorted(glob.glob(os.path.join(D, 'edge_degree_*.pdf'))):
    name = os.path.basename(pdf); lang = 'en' if '_en-' in name else 'zh'
    pt = norm(subprocess.run(['pdftotext', pdf, '-'], capture_output=True, text=True).stdout)
    dd = Document(pdf[:-4] + '.docx'); dt = norm('\n'.join([p.text for p in dd.paragraphs] + [p.text for t in dd.tables for r in t.rows for c in r.cells for p in c.paragraphs]))
    pat = re.compile(r'(?:\d{4}\.\d{2}|[A-Z][a-z]{2} \d{4}) – (?:\d{4}\.\d{2}|[A-Z][a-z]{2} \d{4})(?:（[^）]*）| \([^)]*\))?')
    edu_p = lambda t: t[t.find('教育背景' if lang == 'zh' else 'EDUCATION'):t.find('项目经历' if lang == 'zh' else 'PROJECT EXPERIENCE')]
    a, b = pat.findall(edu_p(pt)), pat.findall(edu_p(dt))
    check(a == b == DEG[lang], f'{name}: 学位状态 PDF={a} DOCX={b}')
    res[name] = {'pdf_degree': a, 'docx_degree': b}

json.dump({'results': res, 'fails': fails}, open(os.path.join(D, 'verify_report.json'), 'w'), ensure_ascii=False, indent=1)
print(json.dumps({'files': len(res), 'fails': fails}, ensure_ascii=False, indent=1))
sys.exit(1 if fails else 0)
